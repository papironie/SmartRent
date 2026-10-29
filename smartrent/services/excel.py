"""Excel in and out with openpyxl.

Export : build a styled .xlsx from rows.
Import : read a .xlsx of properties/units (with optional map coordinates),
         validate EVERY row, and apply all of it or none of it.
"""
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..extensions import db
from ..models import Property, Unit

IMPORT_COLUMNS = ["property", "address", "latitude", "longitude", "unit_no", "monthly_rent"]


# ---------------- export ----------------
def build_workbook(sheet_name: str, header: list[str], rows: list[list]) -> bytes:
    return build_multi_workbook([(sheet_name, header, rows)])


def build_multi_workbook(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    """One styled sheet per (name, header, rows)."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        _write_sheet(wb.create_sheet(name[:31]), header, rows)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_sheet(ws, header: list[str], rows: list[list]) -> None:
    ws.append(header)
    for row in rows:
        ws.append(list(row))
    fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)


def import_template() -> bytes:
    sample = [
        ["Acacia Apartments", "12 Acacia St., Poblacion", 15.9761, 120.5711, "A-01", 8000],
        ["Acacia Apartments", "12 Acacia St., Poblacion", 15.9761, 120.5711, "A-02", 8500],
    ]
    return build_workbook("Units", IMPORT_COLUMNS, sample)


# ---------------- import ----------------
class ImportResult:
    def __init__(self):
        self.properties_created = 0
        self.units_created = 0
        self.units_updated = 0
        self.errors: list[str] = []

    @property
    def ok(self) -> bool:
        return not self.errors


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
    return None if v == "" else v


def import_units(file_obj, owner_id: int) -> ImportResult:
    """Properties are matched by name among this landlord's own properties; new ones belong to them."""
    res = ImportResult()
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        res.errors.append("Could not read the file. Upload a real .xlsx workbook.")
        return res
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        res.errors.append("The sheet is empty.")
        return res
    names = [str(h).strip().lower() if h is not None else "" for h in header]
    missing = [c for c in ("property", "unit_no", "monthly_rent") if c not in names]
    if missing:
        res.errors.append("Missing column(s): " + ", ".join(missing) + ". Download the template to see the layout.")
        return res
    idx = {n: i for i, n in enumerate(names)}

    def cell(row, name):
        i = idx.get(name)
        return _clean(row[i]) if i is not None and i < len(row) else None

    seen: set[tuple[str, str]] = set()
    props: dict[str, Property] = {}
    for line, row in enumerate(rows, start=2):
        if row is None or all(_clean(c) is None for c in row):
            continue
        pname, unit_no = cell(row, "property"), cell(row, "unit_no")
        if not pname or unit_no is None:
            res.errors.append(f"Row {line}: property and unit_no are required.")
            continue
        pname, unit_no = str(pname), str(unit_no)
        try:
            rent = float(cell(row, "monthly_rent"))
            if rent <= 0:
                raise ValueError
        except (TypeError, ValueError):
            res.errors.append(f"Row {line}: monthly_rent must be a number above zero.")
            continue
        lat, lng = cell(row, "latitude"), cell(row, "longitude")
        try:
            lat = float(lat) if lat is not None else None
            lng = float(lng) if lng is not None else None
            if (lat is None) != (lng is None):
                raise ValueError("pair")
            if lat is not None and not (-90 <= lat <= 90 and -180 <= lng <= 180):
                raise ValueError("range")
        except ValueError:
            res.errors.append(f"Row {line}: latitude/longitude must both be given and be valid coordinates.")
            continue
        key = (pname.lower(), unit_no.lower())
        if key in seen:
            res.errors.append(f"Row {line}: unit {unit_no} of {pname} appears twice in the file.")
            continue
        seen.add(key)

        prop = props.get(pname.lower()) or db.session.scalar(
            db.select(Property).where(Property.owner_id == owner_id, db.func.lower(Property.name) == pname.lower())
        )
        address = cell(row, "address")
        if prop is None:
            if not address:
                res.errors.append(f"Row {line}: new property '{pname}' needs an address.")
                continue
            prop = Property(owner_id=owner_id, name=pname, address=str(address))
            db.session.add(prop)
            db.session.flush()
            res.properties_created += 1
        elif address:
            prop.address = str(address)
        if lat is not None:
            prop.set_location(lat, lng)
        props[pname.lower()] = prop

        unit = db.session.scalar(
            db.select(Unit).where(Unit.property_id == prop.id, db.func.lower(Unit.unit_no) == unit_no.lower())
        )
        if unit is None:
            db.session.add(Unit(property_id=prop.id, unit_no=unit_no, monthly_rent=rent))
            res.units_created += 1
        else:
            unit.monthly_rent = rent
            res.units_updated += 1
        db.session.flush()
    wb.close()

    if res.errors:  # all or nothing
        db.session.rollback()
        res.properties_created = res.units_created = res.units_updated = 0
    else:
        db.session.commit()
    return res
