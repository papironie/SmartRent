from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from smartrent.extensions import db
from smartrent.services import analytics, billing
from tests.conftest import login


def test_new_kpis(make):
    p = make.prop()
    make.lease(make.tenant(), make.unit(p, "1", 8000), start=date(2026, 3, 11), end=date(2026, 11, 1))  # ends in 51 days
    make.lease(make.tenant(), make.unit(p, "2", 4000), start=date(2025, 9, 11), end=date(2027, 9, 10))
    make.unit(p, "3", 4000)                                          # vacant
    make.unit(p, "4", 4000, status="maintenance")
    billing.generate(date(2026, 8, 1)); billing.generate(date(2026, 9, 1))
    k = analytics.kpis()
    assert k["economic_occupancy"] == 60.0                           # 12000 / (12000 + 4000 + 4000)
    assert k["avg_rent"] == 6000 and k["billed_change_pct"] == 0.0
    assert (k["at_risk_rent"], k["at_risk_leases"]) == (8000, 1)
    assert 8.5 < k["avg_tenure_months"] < 9.5                        # 6 and 12 months


def test_property_filter(make):
    a = make.prop("A"); b = make.prop("B")
    make.lease(make.tenant(), make.unit(a, "1", 8000))
    make.unit(b, "1", 3000)
    assert analytics.kpis(a.id)["occupancy_rate"] == 100.0
    assert analytics.kpis(b.id)["occupancy_rate"] == 0.0 and analytics.kpis(b.id)["rent_roll"] == 0
    assert analytics.kpis([a.id, b.id])["occupancy_rate"] == 50.0
    assert analytics.kpis([])["units_total"] == 0              # a landlord with no properties sees nothing


def test_occupancy_trend_from_charges(make):
    p = make.prop()
    make.lease(make.tenant(), make.unit(p, "1"), start=date(2026, 8, 1))
    make.unit(p, "2")
    billing.generate(date(2026, 8, 1)); billing.generate(date(2026, 9, 1))
    m = analytics.monthly(12)
    assert m["occupancy"][-2:] == [50.0, 50.0] and m["occupied_units"][-3:] == [0, 1, 1]


def test_forecast_contracted_and_at_risk(make):
    p = make.prop()
    make.lease(make.tenant(), make.unit(p, "1", 1000), end=date(2026, 10, 15))
    make.lease(make.tenant(), make.unit(p, "2", 2000), end=date(2027, 12, 31))
    make.unit(p, "3", 500)
    f = analytics.forecast(3)
    assert f["labels"] == ["Sep 2026", "Oct 2026", "Nov 2026"]
    assert f["contracted"] == [3000, 3000, 2000] and f["at_risk"] == [0, 0, 1000] and f["vacant"] == [500] * 3


def test_tenure_buckets(make):
    p = make.prop()
    make.lease(make.tenant(), make.unit(p, "1"), start=date(2026, 8, 1))
    make.lease(make.tenant(), make.unit(p, "2"), start=date(2024, 1, 1), end=date(2027, 1, 1))
    assert analytics.tenure()["values"] == [1, 0, 0, 1]


def test_unit_performance_and_rent_stats(make):
    p = make.prop("P")
    old = make.lease(make.tenant("Old"), make.unit(p, "1", 4000), start=date(2026, 1, 1))
    billing.generate(date(2026, 5, 1))
    old.status = "ended"; old.unit.status = "vacant"; db.session.commit()
    make.lease(make.tenant("Ana"), make.unit(p, "2", 5000), start=date(2026, 3, 11))
    make.unit(p, "3", 9000)
    rows = {r["unit_no"]: r for r in analytics.records(analytics.unit_performance())}
    assert rows["1"]["idle_months"] == 4 and rows["1"]["last_billed"] == date(2026, 5, 1)
    assert rows["2"]["tenant"] == "Ana" and rows["2"]["tenure_months"] == 6.0
    assert rows["3"]["last_billed"] is None and rows["3"]["idle_months"] is None
    assert rows["3"]["vs_median_pct"] == 80.0                        # median 5000
    s = analytics.rent_stats()[0]
    assert (s["units"], s["median"], s["occupied_avg"], s["vacant_avg"]) == (3, 5000, 5000, 6500)
    text = " ".join(m for _, m in analytics.insights())
    assert "3+ months" in text and "P 3 asks 80.0% more" in text


def test_empty_database_does_not_crash(make):
    assert analytics.forecast()["contracted"] == [0] * 6
    assert analytics.tenure()["values"] == [0, 0, 0, 0]
    assert analytics.unit_performance().empty and analytics.rent_stats() == []
    assert analytics.monthly()["occupancy"] == [0.0] * 12
    assert analytics.kpis()["billed_change_pct"] is None
    assert analytics.insights() == [("ok", "No issues found.")]


def test_analytics_page_api_and_exports(client, make):
    make.admin()
    p = make.prop("Acacia")
    make.lease(make.tenant(), make.unit(p, "1"))
    login(client, "admin@t.test", "admin123")
    assert b"Unit performance" in client.get(f"/admin/analytics?property={p.id}").data
    assert client.get("/admin/analytics?property=nope").status_code == 200
    j = client.get(f"/api/analytics?property={p.id}").get_json()
    assert j["property_id"] == p.id and {"forecast", "tenure"} <= j.keys()
    r = client.get(f"/admin/reports/analytics.xlsx?property={p.id}")
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Monthly", "Forecast", "Lease expiry", "Tenure", "Rent by property", "Units", "Demand", "Insights"]
    assert wb["Summary"]["B2"].value == "Acacia"
    assert client.get("/admin/reports/unit_performance.csv").status_code == 200
