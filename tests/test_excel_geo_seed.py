import io
from datetime import date

from openpyxl import Workbook, load_workbook

from smartrent.extensions import db
from smartrent.models import Property, Unit
from smartrent.services import excel, geo, seed


def _owner() -> int:
    """The default test landlord (same account conftest's factories use)."""
    from smartrent.extensions import db
    from smartrent.models import User
    u = db.session.scalar(db.select(User).where(User.email == "admin@t.test"))
    if u is None:
        u = User(name="Landlord", email="admin@t.test", role="landlord")
        u.set_password("admin123")
        db.session.add(u)
        db.session.commit()
    return u.id
from conftest import login


def xlsx(rows, header=("property", "address", "latitude", "longitude", "unit_no", "monthly_rent")):
    wb = Workbook(); ws = wb.active; ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return b


# ---------- import ----------
def test_import_creates_properties_units_and_pins(ctx):
    r = excel.import_units(owner_id=_owner(), file_obj=xlsx([("Acacia", "12 Acacia", 15.97, 120.57, "A-01", 8000), ("Acacia", None, None, None, "A-02", 8500),
                                 ("Narra", "45 Narra", None, None, "T-01", 12000)]))
    assert r.ok and (r.properties_created, r.units_created, r.units_updated) == (2, 3, 0)
    acacia = db.session.scalar(db.select(Property).where(Property.name == "Acacia"))
    assert round(acacia.lat, 2) == 15.97
    assert db.session.scalar(db.select(db.func.count(Unit.id))) == 3


def test_import_updates_existing_unit_rent(ctx, make):
    p = make.prop("Acacia"); make.unit(p, "A-01", 1000)
    r = excel.import_units(owner_id=_owner(), file_obj=xlsx([("acacia", None, None, None, "a-01", 2222)]))
    assert r.ok and r.units_updated == 1 and r.units_created == 0
    assert float(db.session.scalar(db.select(Unit.monthly_rent))) == 2222.0


def test_import_is_all_or_nothing(ctx):
    r = excel.import_units(owner_id=_owner(), file_obj=xlsx([("Good", "addr", None, None, "1", 5000), ("Bad", "addr", None, None, "1", "abc"),
                                 ("NoAddr", None, None, None, "1", 5000), ("Good", None, None, None, "1", 7000),
                                 ("Pin", "addr", 15.0, None, "1", 5000), ("Pin2", "addr", 99, 120, "1", 5000)]))
    assert not r.ok and len(r.errors) == 5
    assert db.session.scalar(db.select(db.func.count(Property.id))) == 0
    assert db.session.scalar(db.select(db.func.count(Unit.id))) == 0


def test_import_rejects_junk(ctx):
    assert "Could not read" in excel.import_units(io.BytesIO(b"not a workbook"), _owner()).errors[0]
    assert "Missing column" in excel.import_units(owner_id=_owner(), file_obj=xlsx([("a", 1)], header=("property", "foo"))).errors[0]


def test_import_upload_route(client, make):
    make.admin(); login(client, "admin@t.test", "admin123")
    r = client.post("/admin/properties/import", data={"file": (xlsx([("P", "addr", None, None, "1", 100)]), "u.xlsx")},
                    content_type="multipart/form-data", follow_redirects=True)
    assert b"1 unit(s) created" in r.data
    r = client.post("/admin/properties/import", data={"file": (io.BytesIO(b"x"), "u.csv")}, content_type="multipart/form-data",
                    follow_redirects=True)
    assert b"Choose an .xlsx" in r.data
    t = client.get("/admin/properties/import/template.xlsx")
    assert load_workbook(io.BytesIO(t.data)).active["A1"].value == "property"


# ---------- reports ----------
def test_reports_download_formats(client, make):
    p = make.prop("P"); t = make.tenant("Zed"); make.lease(t, make.unit(p, "7"))
    make.admin(); login(client, "admin@t.test", "admin123")
    csvr = client.get("/admin/reports/tenants.csv")
    assert csvr.data.startswith(b"\xef\xbb\xbf") and b"Zed" in csvr.data
    for kind in ("tenants", "leases", "rent_roll", "charges"):
        wb = load_workbook(io.BytesIO(client.get(f"/admin/reports/{kind}.xlsx?month=2026-09").data))
        assert wb.active.max_row >= 1
    assert client.get("/admin/reports/payments.csv").status_code == 404       # payments are gone
    assert client.get("/admin/reports/tenants.pdf").status_code == 404


# ---------- external APIs (mocked: no internet in tests) ----------
class FakeResp:
    def __init__(self, payload, status=200):
        self.payload, self.status = payload, status

    def raise_for_status(self):
        if self.status >= 400:
            import requests
            raise requests.HTTPError("boom")

    def json(self):
        return self.payload


METEO = {"current": {"temperature_2m": 31.5, "relative_humidity_2m": 70, "precipitation": 0, "wind_speed_10m": 9.1, "weather_code": 63},
         "daily": {"precipitation_sum": [1, 2, None, 0, 4, 0, 1, 10, 5, 0]}}


def test_geo_parsers(ctx, monkeypatch):
    def fake_get(url, params=None, timeout=None):
        assert timeout
        if "opentopodata" in url:
            assert params["locations"] == "15.97,120.57"
            return FakeResp({"results": [{"elevation": 54.2}]})
        return FakeResp(METEO)
    monkeypatch.setattr(geo.requests, "get", fake_get)
    geo._cache.clear()
    assert geo.fetch_elevation(15.97, 120.57) == 54.2
    w = geo.fetch_weather(15.97, 120.57)
    assert w["rain_past_7d_mm"] == 8.0 and w["rain_next_3d_mm"] == 15.0 and w["condition"] == "Rain"


def test_geo_failures_return_none(ctx, monkeypatch):
    import requests
    def boom(*a, **k):
        raise requests.ConnectionError("offline")
    monkeypatch.setattr(geo.requests, "get", boom)
    geo._cache.clear()
    assert geo.fetch_elevation(1, 1) is None and geo.fetch_weather(1, 1) is None
    monkeypatch.setattr(geo.requests, "get", lambda *a, **k: FakeResp({"unexpected": 1}))
    assert geo.fetch_elevation(1, 1) is None and geo.fetch_weather(1, 1) is None


def test_site_endpoint_caches_elevation_and_degrades(client, make, monkeypatch):
    p = make.prop("P", 15.97, 120.57); nopin = make.prop("N")
    make.admin(); login(client, "admin@t.test", "admin123")
    calls = {"topo": 0}
    def fake_get(url, params=None, timeout=None):
        if "opentopodata" in url:
            calls["topo"] += 1
            return FakeResp({"results": [{"elevation": 54.2}]})
        return FakeResp(METEO)
    monkeypatch.setattr(geo.requests, "get", fake_get); geo._cache.clear()
    j = client.get(f"/api/properties/{p.id}/site").json
    assert j["elevation_m"] == 54.2 and j["weather"]["temperature_c"] == 31.5 and j["warnings"] == []
    client.get(f"/api/properties/{p.id}/site")
    assert calls["topo"] == 1                                              # elevation stored in PostgreSQL
    assert client.get(f"/api/properties/{nopin.id}/site").status_code == 422
    import requests
    monkeypatch.setattr(geo.requests, "get", lambda *a, **k: (_ for _ in ()).throw(requests.Timeout()))
    geo._cache.clear(); db.session.expire_all()
    j = client.get(f"/api/properties/{p.id}/site").json
    assert j["elevation_m"] == 54.2 and j["weather"] is None and j["warnings"] == ["Weather service unavailable."]


def test_moving_the_pin_clears_cached_elevation(make):
    p = make.prop("P", 15.97, 120.57); p.elevation_m = 50; db.session.commit()
    p.set_location(16.0, 121.0); db.session.commit()
    assert p.elevation_m is None


# ---------- seed ----------
def test_seed_runs_and_is_repeatable(ctx):
    a = seed.run(); b = seed.run()
    assert a == b and a["units"] == 25 and a["tenants"] == 20 and a["charges"] > 100
    from smartrent.services import analytics
    k = analytics.kpis()
    assert k["maintenance"] == 1 and k["active_tenants"] == 20 and k["leases_ending"] == 2
